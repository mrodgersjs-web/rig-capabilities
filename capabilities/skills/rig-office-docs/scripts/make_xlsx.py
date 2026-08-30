#!/usr/bin/env python
"""Render a .xlsx artifact from a JSON spec (clean-room, MIT).

Usage: make_xlsx.py <spec.json>
Spec shape: see SKILL.md "make_xlsx.py spec". Output path is spec["path"],
resolved relative to the directory containing the spec file.
Cell values that are strings starting with '=' are stored as live formulas.
Exit 0 and print the artifact path on success.
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font


def _apply_number_formats(ws, mapping):
    for cell_range, fmt in mapping.items():
        for row in ws[cell_range]:
            for cell in (row if isinstance(row, tuple) else (row,)):
                cell.number_format = fmt


def _build_sheet(wb, sheet_spec, first=False):
    name = sheet_spec.get("name")
    if first:
        ws = wb.active
        if name:
            ws.title = name
    else:
        ws = wb.create_sheet(title=name or None)
    for row in sheet_spec.get("rows", []):
        ws.append(list(row))
    if bool(sheet_spec.get("header_bold", False)) and sheet_spec.get("rows"):
        for cell in ws[1]:
            cell.font = Font(bold=True)
    for cell_ref, value in (sheet_spec.get("cells") or {}).items():
        ws[cell_ref] = value
    for cell_range, fmt in (sheet_spec.get("number_formats") or {}).items():
        _apply_number_formats(ws, {cell_range: fmt})
    for col, width in (sheet_spec.get("widths") or {}).items():
        ws.column_dimensions[col].width = float(width)
    if sheet_spec.get("freeze"):
        ws.freeze_panes = sheet_spec["freeze"]
    return ws


def build(spec, out_path):
    wb = Workbook()
    sheets = spec.get("sheets") or [{"name": "Sheet"}]
    for idx, sheet_spec in enumerate(sheets):
        _build_sheet(wb, sheet_spec, first=(idx == 0))
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_path))
    return out_path


def main(argv):
    if len(argv) != 2:
        print(__doc__, file=sys.stderr)
        return 2
    spec_path = Path(argv[1]).resolve()
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    out = Path(spec["path"])
    if not out.is_absolute():
        out = spec_path.parent / out
    built = build(spec, out)
    print(f"xlsx written: {built} ({built.stat().st_size} bytes)")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
