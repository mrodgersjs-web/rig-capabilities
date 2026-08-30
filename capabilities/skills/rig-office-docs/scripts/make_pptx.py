#!/usr/bin/env python
"""Render a .pptx artifact from a JSON spec (clean-room, MIT).

Usage: make_pptx.py <spec.json>
Spec shape: see SKILL.md "make_pptx.py spec". Output path is spec["path"],
resolved relative to the directory containing the spec file.
Exit 0 and print the artifact path on success.
"""
import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches


def _write_bullets(text_frame, bullets):
    """bullets: list of str or {"text": str, "level": int}."""
    text_frame.clear()
    for i, item in enumerate(bullets):
        if isinstance(item, dict):
            text, level = str(item.get("text", "")), int(item.get("level", 0))
        else:
            text, level = str(item), 0
        para = text_frame.paragraphs[0] if i == 0 else text_frame.add_paragraph()
        para.text = text
        para.level = max(0, level)


def _add_table(slide, spec_el):
    rows = spec_el.get("rows") or [[]]
    cols = max(len(r) for r in rows)
    left = Inches(float(spec_el.get("x_in", 0.75)))
    top = Inches(float(spec_el.get("y_in", 1.6)))
    width = Inches(float(spec_el.get("width_in", 8.0)))
    height = Inches(float(spec_el.get("height_in", 0.4 * len(rows) + 0.2)))
    shape = slide.shapes.add_table(len(rows), cols, left, top, width, height)
    table = shape.table
    for i, row in enumerate(rows):
        for j in range(cols):
            value = row[j] if j < len(row) else ""
            cell = table.cell(i, j)
            cell.text = "" if value is None else str(value)
            if bool(spec_el.get("header", True)) and i == 0:
                for p in cell.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.bold = True
    return shape


def _build_slide(prs, slide_spec):
    title = slide_spec.get("title")
    layout_index = slide_spec.get("layout")
    if layout_index is None:
        has_body = any(k in slide_spec for k in ("bullets", "text", "table"))
        has_sub = "subtitle" in slide_spec
        layout_index = 0 if (title and has_sub and not has_body) else (1 if title else 6)
    slide = prs.slides.add_slide(prs.slide_layouts[int(layout_index)])

    if title is not None:
        for ph in slide.placeholders:
            if ph.placeholder_format.idx == 0:
                ph.text = str(title)
                break

    body_ph = None
    for ph in slide.placeholders:
        if ph.placeholder_format.idx != 0 and ph != slide.shapes.title:
            body_ph = ph
            break

    if "subtitle" in slide_spec and body_ph is not None:
        body_ph.text = str(slide_spec["subtitle"])
    if "bullets" in slide_spec:
        if body_ph is None:
            box = slide.shapes.add_textbox(Inches(0.75), Inches(1.6), Inches(8.5), Inches(5.0))
            body_ph = box
        _write_bullets(body_ph.text_frame, slide_spec["bullets"])
    elif "text" in slide_spec and body_ph is not None:
        body_ph.text_frame.text = str(slide_spec["text"])

    for el in slide_spec.get("elements", []):
        if el.get("type") == "table":
            _add_table(slide, el)
        else:
            raise ValueError(f"unknown slide element type: {el.get('type')!r}")

    if slide_spec.get("notes"):
        slide.notes_slide.notes_text_frame.text = str(slide_spec["notes"])
    return slide


def build(spec, out_path):
    prs = Presentation(spec["template"]) if spec.get("template") else Presentation()
    for slide_spec in spec.get("slides", []):
        _build_slide(prs, slide_spec)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return out_path


def main(argv):
    if len(argv) != 2:
        print(__doc__, file=sys.stderr)
        return 2
    spec_path = Path(argv[1]).resolve()
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    out = Path(spec["path"])
    if not out.is_absolute():
        out = spec_path.parent / out
    built = build(spec, out)
    print(f"pptx written: {built} ({built.stat().st_size} bytes)")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
